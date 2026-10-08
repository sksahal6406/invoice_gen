import io
from decimal import Decimal, ROUND_HALF_UP
from datetime import date

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


# ── Fixed seller / template data ──

SELLER = {
    "name": "DIAMOND PHARMA DISTRIBUTORS",
    "address": "A.32, PATEL SUPER MARKET, STATION ROAD, BHARUCH-392001",
    "phone": "(P)02642-267181, (M) 98243 91122 / 9624488166",
    "gstin": "24ABRPC7309Q1Z5",
    "bank_ac": "10082001001347",
    "ifsc": "PMEC0100809",
    "bank_name": "PRIME CO-OP.BANK LTD",
    "gpay": "9824391122",
    "dl_no": "20B/BHA/184713, 21B/BHA/184714, 20D/BHA/184715",
    "software": "VISUAL INFOSOFT PVT. LTD.",
    "customer_care": "079 3520 7999",
}

FIXED_NOTES = [
    "* GOODS ONCE SOLD WILL NOT BE TAKEN BACK.",
    "* ALL DISPUTES SUBJECT TO BHARUCH JURISDICTION ONLY.",
    "* SALES MADE TO YOUR CAPACITY AS RETAILER.",
]

FIXED_MESSAGE = "Message: *** ALL VACCINES ARE AVAILABLE***"

SGST_RATE = Decimal("2.50")
CGST_RATE = Decimal("2.50")

PRINT_ROWS = 8


# ── Helpers ──

def D(value):
    if value is None or value == "":
        return Decimal("0")
    try:
        f = float(value)
        if f != f:
            return Decimal("0")
        return Decimal(str(value).strip())
    except Exception:
        return Decimal("0")


def money(value):
    return D(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def amount_in_words(n):
    n = int(Decimal(str(n)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    ones = [
        "Zero", "One", "Two", "Three", "Four", "Five", "Six", "Seven",
        "Eight", "Nine", "Ten", "Eleven", "Twelve", "Thirteen", "Fourteen",
        "Fifteen", "Sixteen", "Seventeen", "Eighteen", "Nineteen",
    ]
    tens = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty",
            "Seventy", "Eighty", "Ninety"]

    def under_1000(x):
        parts = []
        if x >= 100:
            parts.append(ones[x // 100] + " Hundred")
            x %= 100
            if x:
                parts.append("and")
        if x >= 20:
            parts.append(tens[x // 10])
            if x % 10:
                parts.append(ones[x % 10])
        elif x > 0:
            parts.append(ones[x])
        return " ".join(parts)

    if n == 0:
        return "Zero Only"
    parts = []
    crore = n // 10_000_000
    n %= 10_000_000
    lakh = n // 100_000
    n %= 100_000
    thousand = n // 1_000
    n %= 1_000
    if crore:
        parts.append(under_1000(crore) + " Crore")
    if lakh:
        parts.append(under_1000(lakh) + " Lakh")
    if thousand:
        parts.append(under_1000(thousand) + " Thousand")
    if n:
        parts.append(under_1000(n))
    return " ".join(parts) + " Only"


def calculate_products(rows):
    calculated = []
    t_tax = t_sgst = t_cgst = t_amt = t_qty = Decimal("0")
    sr = 1
    for r in rows:
        desc = str(r.get("description", "")).strip()
        hsn = str(r.get("hsn", "")).strip()
        qty = D(r.get("qty"))
        rate = D(r.get("rate"))
        discount = D(r.get("discount"))
        if not desc and not hsn and not qty and not rate:
            continue
        gross = money(qty * rate)
        disc_amt = money(gross * discount / Decimal("100"))
        taxable = money(gross - disc_amt)
        sgst = money(taxable * SGST_RATE / Decimal("100"))
        cgst = money(taxable * CGST_RATE / Decimal("100"))
        amount = money(taxable + sgst + cgst)
        calculated.append({
            **r, "sr": sr,
            "taxable": taxable, "sgst": sgst, "cgst": cgst, "amount": amount,
        })
        t_tax += taxable
        t_sgst += sgst
        t_cgst += cgst
        t_amt += amount
        t_qty += qty
        sr += 1
    return (calculated, money(t_tax), money(t_sgst),
            money(t_cgst), money(t_amt), t_qty)


# ── PDF helpers ──

def dt(c, text, x, y, size=7, bold=False, align="left"):
    font = "Helvetica-Bold" if bold else "Helvetica"
    c.setFont(font, size)
    s = str(text)
    if align == "right":
        c.drawRightString(x, y, s)
    elif align == "center":
        c.drawCentredString(x, y, s)
    else:
        c.drawString(x, y, s)


def draw_wrapped(c, text, x, y, max_w, size=7, leading=9, bold=False, max_lines=4):
    font = "Helvetica-Bold" if bold else "Helvetica"
    c.setFont(font, size)
    words = str(text).split()
    lines, cur = [], ""
    for w in words:
        test = w if not cur else cur + " " + w
        if stringWidth(test, font, size) <= max_w:
            cur = test
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    for i, line in enumerate(lines[:max_lines]):
        c.drawString(x, y - i * leading, line)


# ══════════════════════════════════════════════
#  PDF GENERATION
# ══════════════════════════════════════════════

def draw_pdf(inv, prods):
    buf = io.BytesIO()
    pw, ph = landscape(A4)  # 841.89 x 595.28
    c = canvas.Canvas(buf, pagesize=(pw, ph))

    L, B = 28, 24
    R, T = pw - 28, ph - 24
    W = R - L

    # ── Outer border ──
    c.setLineWidth(0.8)
    c.rect(L, B, W, T - B)

    # ────────────────────────────────────
    #  HEADER  (T to T-118)
    # ────────────────────────────────────
    hdr_btm = T - 118
    c.line(L, hdr_btm, R, hdr_btm)

    # Two vertical dividers creating 3 header zones
    div1 = L + 468          # seller | metadata
    div2 = div1 + 127       # metadata | customer
    c.setLineWidth(0.5)
    c.line(div1, T, div1, hdr_btm)
    c.line(div2, T, div2, hdr_btm)

    # ── Zone 1: Seller ──
    x0 = L + 7
    dt(c, SELLER["name"], x0, T - 16, 12, True)
    dt(c, SELLER["address"], x0, T - 30, 7)
    dt(c, SELLER["phone"], x0, T - 41, 7)
    dt(c, FIXED_NOTES[0], x0, T - 58, 7)
    dt(c, "GST No: " + SELLER["gstin"], x0, T - 70, 7)

    # ── Zone 2: Invoice number box + metadata ──
    box_l = div1 + 3
    box_r = div2 - 3
    box_t = T - 3
    box_b = T - 30
    c.rect(box_l, box_b, box_r - box_l, box_t - box_b)
    dt(c, inv["invoice_no"], (box_l + box_r) / 2, box_b + 7, 14, True, "center")

    mx = div1 + 8
    dt(c, "Inv Date", mx, T - 44, 7)
    dt(c, ": " + inv["inv_date"], mx + 42, T - 44, 7)
    dt(c, "Mode", mx, T - 57, 7)
    dt(c, ": " + inv["mode"], mx + 42, T - 57, 7)
    dt(c, "Due Date", mx, T - 70, 7)
    dt(c, ": " + inv["due_date"], mx + 42, T - 70, 7)

    # ── Zone 3: TAX INVOICE + Customer ──
    dt(c, "TAX INVOICE", R - 8, T - 14, 10, True, "right")

    cx = div2 + 6
    cw = R - div2 - 12
    dt(c, "To, " + inv["customer_name"], cx, T - 30, 7, True)
    draw_wrapped(c, inv["customer_address"], cx, T - 42, cw, 7, 9, False, 3)
    dt(c, "D.L.No:", cx, T - 80, 7)
    dt(c, inv["customer_dl"], cx + 36, T - 80, 7)
    dt(c, "GSTIN: " + inv["customer_gstin"], cx, T - 93, 7)
    dt(c, "PAN: " + inv["customer_pan"], cx + 120, T - 93, 7)

    # ────────────────────────────────────
    #  PRODUCT TABLE  (hdr_btm to tbl_btm)
    # ────────────────────────────────────
    tbl_top = hdr_btm
    tbl_btm = 168

    col_defs = [
        ("Sr.", 23), ("HSN Code", 47), ("Description of Goods", 135),
        ("Batch", 63), ("Exp", 42), ("MRP", 50), ("Qty/Fr", 42),
        ("Rate", 50), ("Dis.%", 42), ("Taxable", 60), ("(%)", 28),
        ("SGST", 50), ("(%)", 28), ("CGST", 50), ("Amount", 76),
    ]
    raw_w = sum(w for _, w in col_defs)
    sc = W / raw_w
    widths = [w * sc for _, w in col_defs]

    xp = [L]
    for w in widths:
        xp.append(xp[-1] + w)

    # Column header row
    hdr_h = 22
    c.setLineWidth(0.5)
    c.rect(L, tbl_top - hdr_h, W, hdr_h)
    for x in xp[1:-1]:
        c.line(x, tbl_top, x, tbl_btm)

    for i, (label, _) in enumerate(col_defs):
        mid = (xp[i] + xp[i + 1]) / 2
        dt(c, label, mid, tbl_top - 15, 6, True, "center")

    # 9 content rows (8 product + 1 totals)
    n_rows = PRINT_ROWS + 1
    row_h = (tbl_top - hdr_h - tbl_btm) / n_rows

    for r in range(n_rows + 1):
        y = tbl_top - hdr_h - r * row_h
        c.line(L, y, R, y)

    # Product data rows
    for i in range(PRINT_ROWS):
        y = tbl_top - hdr_h - (i + 0.5) * row_h
        if i < len(prods):
            p = prods[i]
            vals = [
                str(p["sr"]),
                str(p.get("hsn", "")),
                str(p.get("description", "")),
                str(p.get("batch", "")),
                str(p.get("exp", "")),
                f'{D(p.get("mrp", 0)):.2f}' if D(p.get("mrp", 0)) else "",
                f'{D(p.get("qty", 0)):g}' if D(p.get("qty", 0)) else "",
                f'{D(p.get("rate", 0)):.2f}' if D(p.get("rate", 0)) else "",
                f'{D(p.get("discount", 0)):.2f}' if D(p.get("discount", 0)) else "",
                f'{p["taxable"]:.2f}',
                f'{SGST_RATE:.2f}',
                f'{p["sgst"]:.2f}',
                f'{CGST_RATE:.2f}',
                f'{p["cgst"]:.2f}',
                f'{p["amount"]:.2f}',
            ]
        else:
            vals = [""] * 15

        for j, v in enumerate(vals):
            if not v:
                continue
            mid = (xp[j] + xp[j + 1]) / 2
            if j == 2:
                dt(c, v, xp[j] + 3, y, 6.5)
            elif j in (5, 7, 8, 9, 11, 13, 14):
                dt(c, v, xp[j + 1] - 3, y, 6.5, False, "right")
            else:
                dt(c, v, mid, y, 6.5, False, "center")

    # ── Totals row (row 9) ──
    totals_row_top = tbl_top - hdr_h - PRINT_ROWS * row_h
    ty = totals_row_top - row_h * 0.5

    dt(c, FIXED_MESSAGE, L + 3, ty, 7, True)
    dt(c, f'{inv["total_qty"]:g}', xp[7] - 3, ty, 6.5, True, "right")
    dt(c, f'{inv["taxable_total"]:.2f}', xp[10] - 3, ty, 6.5, True, "right")
    dt(c, f'{inv["sgst_total"]:.2f}', xp[12] - 3, ty, 6.5, True, "right")
    dt(c, f'{inv["cgst_total"]:.2f}', xp[14] - 3, ty, 6.5, True, "right")
    dt(c, f'{inv["gross_total"]:.2f}', xp[15] - 3, ty, 6.5, True, "right")

    # ────────────────────────────────────
    #  BOTTOM SECTION  (tbl_btm to B)
    #  Height = 168 - 24 = 144 pt
    # ────────────────────────────────────
    bot_top = tbl_btm
    rbox_x = R - 185

    # Right totals box — vertical divider
    c.line(rbox_x, B, rbox_x, bot_top)

    # Row 1: Amount in Words + Balance  (y = 156)
    y1 = bot_top - 12
    dt(c, "Amount in Words: " + inv["amount_words"], L + 4, y1, 7)
    dt(c, "Balance:", rbox_x + 4, y1, 7)
    dt(c, f'{inv["balance"]:.2f}', R - 5, y1, 7, False, "right")

    # Horizontal line below Amount in Words (left side only)
    hl1 = bot_top - 20
    c.line(L, hl1, rbox_x, hl1)

    # Row 2: Note 1 + GST headers + Other +/-  (y = 138)
    y2 = hl1 - 10
    dt(c, FIXED_NOTES[1], L + 4, y2, 7)
    dt(c, "Other +/-", rbox_x + 4, y2, 7)
    dt(c, f'{inv["other"]:.2f}', R - 5, y2, 7, False, "right")

    # Row 3: Note 2 + GST values + Credit Note  (y = 125)
    y3 = y2 - 13
    dt(c, FIXED_NOTES[2], L + 4, y3, 7)
    dt(c, "Credit Note", rbox_x + 4, y3, 7)
    dt(c, f'{inv["credit_note"]:.2f}', R - 5, y3, 7, False, "right")

    # Row 4: Round Off  (y = 112)
    y4 = y3 - 13
    dt(c, "Round Off", rbox_x + 4, y4, 7)
    dt(c, f'{inv["round_off"]:.2f}', R - 5, y4, 7, False, "right")

    # Line above Net Amount
    nl = y4 - 7
    c.line(rbox_x + 2, nl, R - 2, nl)

    # Net Amount  (y = 95)
    y5 = nl - 10
    dt(c, "Net Amount", rbox_x + 4, y5, 8, True)
    dt(c, f'{inv["net_amount"]:.2f}', R - 5, y5, 8, True, "right")

    # For, SELLER  (y = 80)
    dt(c, "For, " + SELLER["name"], rbox_x + 4, y5 - 15, 7, True)

    # Authorized Signatory  (y = 32)
    dt(c, "Authorized Signatory", R - 8, B + 8, 7, False, "right")

    # ── GST Summary (middle zone, same rows as notes) ──
    gst_x = rbox_x - 195
    dt(c, "GST %", gst_x, y2, 7, True)
    dt(c, "Taxable", gst_x + 50, y2, 7, True)
    dt(c, "SGST", gst_x + 105, y2, 7, True)
    dt(c, "CGST", gst_x + 148, y2, 7, True)
    dt(c, "IGST", gst_x + 190, y2, 7, True)

    dt(c, f"{SGST_RATE + CGST_RATE:.2f}", gst_x, y3, 7)
    dt(c, f'{inv["taxable_total"]:.2f}', gst_x + 50, y3, 7)
    dt(c, f'{inv["sgst_total"]:.2f}', gst_x + 105, y3, 7)
    dt(c, f'{inv["cgst_total"]:.2f}', gst_x + 148, y3, 7)
    dt(c, "0.00", gst_x + 190, y3, 7)

    # ── Bank details (left column, below notes) ──
    bank_y = y3 - 18
    bank_sp = 12
    dt(c, "BANK A/C NO-" + SELLER["bank_ac"], L + 4, bank_y, 7)
    dt(c, "IFSC CODE -" + SELLER["ifsc"], L + 4, bank_y - bank_sp, 7)
    dt(c, "BANK NAME- " + SELLER["bank_name"], L + 4, bank_y - 2 * bank_sp, 7)
    dt(c, "G PAY NO=" + SELLER["gpay"], L + 4, bank_y - 3 * bank_sp, 7)
    dt(c, "D.L.NO:" + SELLER["dl_no"], L + 4, bank_y - 4 * bank_sp, 7)

    # ── Other / S.Man / D.Man box (middle zone) ──
    oth_x = gst_x - 5
    oth_top = bank_y - 2 * bank_sp
    c.rect(oth_x, oth_top - 32, 130, 36)
    dt(c, "Other  :.", oth_x + 4, oth_top - 4, 7)
    dt(c, "S.Man :", oth_x + 4, oth_top - 16, 7)
    dt(c, "D.Man :", oth_x + 4, oth_top - 28, 7)

    # ── E. & O. E. ──
    dt(c, "E. & O. E.", rbox_x - 55, B + 8, 7)

    # ── Footer ──
    dt(c, "Software by " + SELLER["software"]
       + " : Customer Care No: " + SELLER["customer_care"],
       L + 4, B + 5, 6.5)

    c.showPage()
    c.save()
    buf.seek(0)
    return buf.getvalue()


# ══════════════════════════════════════════════
#  EXCEL GENERATION
# ══════════════════════════════════════════════

def draw_excel(inv, prods):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tax Invoice"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(
        left=0.2, right=0.2, top=0.25, bottom=0.25, header=0, footer=0)
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False

    thin = Side(style="thin", color="000000")
    bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
    la = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ra = Alignment(horizontal="right", vertical="center")
    f8 = Font(name="Arial", size=8)
    f8b = Font(name="Arial", size=8, bold=True)

    for i, w in enumerate(
            [5, 9, 24, 11, 8, 9, 9, 9, 8, 11, 6, 10, 6, 10, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Header ──
    ws.merge_cells("A1:F1")
    ws["A1"] = SELLER["name"]
    ws["A1"].font = Font(name="Arial", size=16, bold=True)
    ws["A1"].alignment = la

    for rn, txt in [(2, SELLER["address"]), (3, SELLER["phone"]),
                    (4, FIXED_NOTES[0]), (5, "GST No: " + SELLER["gstin"])]:
        ws.merge_cells(start_row=rn, start_column=1, end_row=rn, end_column=6)
        ws.cell(rn, 1, txt).font = Font(name="Arial", size=9)

    ws.merge_cells("N1:O1")
    ws["N1"] = "TAX INVOICE"
    ws["N1"].font = Font(name="Arial", size=12, bold=True)
    ws["N1"].alignment = ra

    ws.merge_cells("G1:I1")
    ws["G1"] = inv["invoice_no"]
    ws["G1"].font = Font(name="Arial", size=14, bold=True)
    ws["G1"].alignment = ctr
    for cc in range(7, 10):
        ws.cell(1, cc).border = bdr

    for cl, lb, cv, vl in [
        ("G3", "Inv Date", "H3", ": " + inv["inv_date"]),
        ("G4", "Mode", "H4", ": " + inv["mode"]),
        ("G5", "Due Date", "H5", ": " + inv["due_date"]),
    ]:
        ws[cl] = lb
        ws[cl].font = f8
        ws[cv] = vl
        ws[cv].font = f8

    ws.merge_cells("J2:O2")
    ws["J2"] = "To, " + inv["customer_name"]
    ws["J2"].font = f8b
    ws["J2"].alignment = la

    ws.merge_cells("J3:O4")
    ws["J3"] = inv["customer_address"]
    ws["J3"].font = f8
    ws["J3"].alignment = la

    ws.merge_cells("J5:L5")
    ws["J5"] = "D.L.No: " + inv["customer_dl"]
    ws["J5"].font = f8

    ws.merge_cells("J6:L6")
    ws["J6"] = "GSTIN: " + inv["customer_gstin"]
    ws["J6"].font = f8

    ws.merge_cells("M6:O6")
    ws["M6"] = "PAN: " + inv["customer_pan"]
    ws["M6"].font = f8

    # ── Product table header (row 8) ──
    headers = [
        "Sr.", "HSN Code", "Description of Goods", "Batch", "Exp", "MRP",
        "Qty/Fr", "Rate", "Dis.%", "Taxable", "(%)", "SGST", "(%)", "CGST",
        "Amount",
    ]
    hr = 8
    for col, h in enumerate(headers, 1):
        cell = ws.cell(hr, col, h)
        cell.font = f8b
        cell.alignment = ctr
        cell.border = bdr

    # ── Product rows (rows 9-16) ──
    fpr = hr + 1
    for i in range(PRINT_ROWS):
        r = fpr + i
        p = prods[i] if i < len(prods) else None
        vals = [""] * 15
        if p:
            vals = [
                p["sr"], p.get("hsn", ""), p.get("description", ""),
                p.get("batch", ""), p.get("exp", ""),
                float(D(p.get("mrp", 0))) if D(p.get("mrp", 0)) else "",
                float(D(p.get("qty", 0))) if D(p.get("qty", 0)) else "",
                float(D(p.get("rate", 0))) if D(p.get("rate", 0)) else "",
                float(D(p.get("discount", 0))) if D(p.get("discount", 0)) else "",
                float(p["taxable"]), float(SGST_RATE),
                float(p["sgst"]), float(CGST_RATE),
                float(p["cgst"]), float(p["amount"]),
            ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(r, col, val)
            cell.font = f8
            cell.border = bdr
            cell.alignment = (la if col == 3
                              else ra if col in (6, 8, 9, 10, 12, 14, 15)
                              else ctr)
            if (col in (6, 8, 9, 10, 11, 12, 13, 14, 15)
                    and isinstance(val, (float, int))):
                cell.number_format = "0.00"

    # ── Totals row (row 17) ──
    tr = fpr + PRINT_ROWS
    ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=6)
    ws.cell(tr, 1, FIXED_MESSAGE).font = f8b
    ws.cell(tr, 1).alignment = la
    ws.cell(tr, 7, float(inv["total_qty"])).font = f8b
    ws.cell(tr, 10, float(inv["taxable_total"])).font = f8b
    ws.cell(tr, 12, float(inv["sgst_total"])).font = f8b
    ws.cell(tr, 14, float(inv["cgst_total"])).font = f8b
    ws.cell(tr, 15, float(inv["gross_total"])).font = f8b
    for col in range(1, 16):
        ws.cell(tr, col).border = bdr
        if col in (7, 10, 12, 14, 15):
            ws.cell(tr, col).number_format = "0.00"
            ws.cell(tr, col).alignment = ra

    # ── Bottom section ──
    bs = tr + 1  # bottom-start row

    # Row bs+0: Amount in Words + Balance
    ws.merge_cells(start_row=bs, start_column=1, end_row=bs, end_column=11)
    ws.cell(bs, 1, "Amount in Words: " + inv["amount_words"]).font = f8
    ws.merge_cells(start_row=bs, start_column=12, end_row=bs, end_column=14)
    ws.cell(bs, 12, "Balance:").font = f8
    ws.cell(bs, 12).alignment = la
    ws.cell(bs, 15, float(inv["balance"])).font = f8
    ws.cell(bs, 15).number_format = "0.00"
    ws.cell(bs, 15).alignment = ra
    for cc in range(12, 16):
        ws.cell(bs, cc).border = bdr

    # Row bs+1: Note 1 + GST headers + Other +/-
    ws.merge_cells(start_row=bs + 1, start_column=1,
                   end_row=bs + 1, end_column=5)
    ws.cell(bs + 1, 1, FIXED_NOTES[1]).font = f8
    for col, lbl in [(6, "GST %"), (7, "Taxable"), (8, "SGST"),
                     (9, "CGST"), (10, "IGST")]:
        ws.cell(bs + 1, col, lbl).font = f8b
        ws.cell(bs + 1, col).alignment = ctr
    ws.merge_cells(start_row=bs + 1, start_column=12,
                   end_row=bs + 1, end_column=14)
    ws.cell(bs + 1, 12, "Other +/-").font = f8
    ws.cell(bs + 1, 12).alignment = la
    ws.cell(bs + 1, 15, float(inv["other"])).font = f8
    ws.cell(bs + 1, 15).number_format = "0.00"
    ws.cell(bs + 1, 15).alignment = ra
    for cc in range(12, 16):
        ws.cell(bs + 1, cc).border = bdr

    # Row bs+2: Note 2 + GST values + Credit Note
    ws.merge_cells(start_row=bs + 2, start_column=1,
                   end_row=bs + 2, end_column=5)
    ws.cell(bs + 2, 1, FIXED_NOTES[2]).font = f8
    for col, val in [(6, "5.00"), (7, float(inv["taxable_total"])),
                     (8, float(inv["sgst_total"])),
                     (9, float(inv["cgst_total"])), (10, 0.00)]:
        ws.cell(bs + 2, col, val).font = f8
        ws.cell(bs + 2, col).alignment = ctr
        if isinstance(val, float):
            ws.cell(bs + 2, col).number_format = "0.00"
    ws.merge_cells(start_row=bs + 2, start_column=12,
                   end_row=bs + 2, end_column=14)
    ws.cell(bs + 2, 12, "Credit Note").font = f8
    ws.cell(bs + 2, 12).alignment = la
    ws.cell(bs + 2, 15, float(inv["credit_note"])).font = f8
    ws.cell(bs + 2, 15).number_format = "0.00"
    ws.cell(bs + 2, 15).alignment = ra
    for cc in range(12, 16):
        ws.cell(bs + 2, cc).border = bdr

    # Row bs+3: Bank 1 + Round Off
    ws.merge_cells(start_row=bs + 3, start_column=1,
                   end_row=bs + 3, end_column=5)
    ws.cell(bs + 3, 1, "BANK A/C NO-" + SELLER["bank_ac"]).font = f8
    ws.merge_cells(start_row=bs + 3, start_column=12,
                   end_row=bs + 3, end_column=14)
    ws.cell(bs + 3, 12, "Round Off").font = f8
    ws.cell(bs + 3, 12).alignment = la
    ws.cell(bs + 3, 15, float(inv["round_off"])).font = f8
    ws.cell(bs + 3, 15).number_format = "0.00"
    ws.cell(bs + 3, 15).alignment = ra
    for cc in range(12, 16):
        ws.cell(bs + 3, cc).border = bdr

    # Row bs+4: Bank 2 + Net Amount
    ws.merge_cells(start_row=bs + 4, start_column=1,
                   end_row=bs + 4, end_column=5)
    ws.cell(bs + 4, 1, "IFSC CODE -" + SELLER["ifsc"]).font = f8
    ws.merge_cells(start_row=bs + 4, start_column=12,
                   end_row=bs + 4, end_column=14)
    ws.cell(bs + 4, 12, "Net Amount").font = f8b
    ws.cell(bs + 4, 12).alignment = la
    ws.cell(bs + 4, 15, float(inv["net_amount"])).font = f8b
    ws.cell(bs + 4, 15).number_format = "0.00"
    ws.cell(bs + 4, 15).alignment = ra
    for cc in range(12, 16):
        ws.cell(bs + 4, cc).border = bdr

    # Row bs+5: Bank 3 + For, SELLER
    ws.merge_cells(start_row=bs + 5, start_column=1,
                   end_row=bs + 5, end_column=5)
    ws.cell(bs + 5, 1, "BANK NAME- " + SELLER["bank_name"]).font = f8
    ws.merge_cells(start_row=bs + 5, start_column=12,
                   end_row=bs + 5, end_column=15)
    ws.cell(bs + 5, 12, "For, " + SELLER["name"]).font = f8b

    # Row bs+6: Bank 4
    ws.merge_cells(start_row=bs + 6, start_column=1,
                   end_row=bs + 6, end_column=5)
    ws.cell(bs + 6, 1, "G PAY NO=" + SELLER["gpay"]).font = f8

    # Row bs+7: Bank 5
    ws.merge_cells(start_row=bs + 7, start_column=1,
                   end_row=bs + 7, end_column=5)
    ws.cell(bs + 7, 1, "D.L.NO:" + SELLER["dl_no"]).font = f8

    # Authorized Signatory
    ws.merge_cells(start_row=bs + 9, start_column=13,
                   end_row=bs + 9, end_column=15)
    ws.cell(bs + 9, 13, "Authorized Signatory").font = f8
    ws.cell(bs + 9, 13).alignment = ra

    # Footer + E. & O. E.
    footer_row = bs + 10
    ws.merge_cells(start_row=footer_row, start_column=1,
                   end_row=footer_row, end_column=9)
    ws.cell(footer_row, 1,
            "Software by " + SELLER["software"]
            + " : Customer Care No: " + SELLER["customer_care"]
            ).font = Font(name="Arial", size=7)
    ws.merge_cells(start_row=footer_row, start_column=10,
                   end_row=footer_row, end_column=11)
    ws.cell(footer_row, 10, "E. & O. E.").font = f8

    ws.print_area = f"A1:O{footer_row}"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


# ══════════════════════════════════════════════
#  STREAMLIT UI
# ══════════════════════════════════════════════

st.set_page_config(page_title="Diamond Pharma Tax Invoice", layout="wide")
st.title("Diamond Pharma Distributors — Tax Invoice")

if "last_invoice_num" not in st.session_state:
    st.session_state.last_invoice_num = 1262
if "num_product_rows" not in st.session_state:
    st.session_state.num_product_rows = 1

# Initialize all product row keys in session state
PRODUCT_KEYS = ["hsn", "description", "batch", "exp",
                "mrp", "qty", "rate", "discount"]
for i in range(PRINT_ROWS):
    for k in PRODUCT_KEYS:
        sk = f"row_{i}_{k}"
        if sk not in st.session_state:
            st.session_state[sk] = ""

# ── Invoice details ──
c1, c2, c3, c4 = st.columns(4)
with c1:
    invoice_no = st.text_input(
        "Invoice No.", value=f"T{st.session_state.last_invoice_num + 1}")
with c2:
    inv_date = st.date_input("Invoice Date", value=date.today())
with c3:
    mode = st.selectbox("Mode", ["CREDIT", "CASH", "UPI", "OTHER"])
with c4:
    due_date = st.date_input("Due Date", value=inv_date)

# ── Customer ──
st.subheader("Customer")
cl, cr = st.columns(2)
with cl:
    customer_name = st.text_input("Customer Name")
    customer_address = st.text_area("Customer Address", height=80)
with cr:
    customer_dl = st.text_input("D.L. No.")
    customer_gstin = st.text_input("GSTIN")
    customer_pan = st.text_input("PAN")

# ── Products ──
st.subheader("Products")

header_labels = ["HSN", "Description", "Batch", "Exp",
                 "MRP", "Qty", "Rate", "Dis.%"]
col_widths = [0.8, 2.2, 1.0, 0.7, 0.9, 0.7, 0.9, 0.7]

hdr_cols = st.columns(col_widths)
for col, lbl in zip(hdr_cols, header_labels):
    with col:
        st.markdown(f"**{lbl}**")

edited_rows = []
for i in range(st.session_state.num_product_rows):
    row_cols = st.columns(col_widths)
    values = {}
    for col, key in zip(row_cols, PRODUCT_KEYS):
        with col:
            values[key] = st.text_input(
                key, key=f"row_{i}_{key}", label_visibility="collapsed")
    edited_rows.append(values)

btn_cols = st.columns([1, 1, 6])
with btn_cols[0]:
    if st.session_state.num_product_rows < PRINT_ROWS:
        if st.button("+ Add Row"):
            st.session_state.num_product_rows += 1
            st.rerun()
with btn_cols[1]:
    if st.session_state.num_product_rows > 1:
        if st.button("- Remove Row"):
            n = st.session_state.num_product_rows - 1
            for k in PRODUCT_KEYS:
                st.session_state[f"row_{n}_{k}"] = ""
            st.session_state.num_product_rows = n
            st.rerun()

products, taxable_total, sgst_total, cgst_total, gross_total, total_qty = \
    calculate_products(edited_rows)

# ── Adjustments ──
st.subheader("Adjustments")
a1, a2 = st.columns(2)
with a1:
    other = st.number_input("Other +/-", value=0.00, step=0.01, format="%.2f")
with a2:
    credit_note = st.number_input(
        "Credit Note", value=0.00, step=0.01, format="%.2f")

gross_adj = money(gross_total + D(other) - D(credit_note))
round_off = money(
    gross_adj.quantize(Decimal("1"), rounding=ROUND_HALF_UP) - gross_adj)
net_amount = money(gross_adj + round_off)

inv_data = {
    "invoice_no": invoice_no,
    "inv_date": inv_date.strftime("%d/%m/%Y"),
    "due_date": due_date.strftime("%d/%m/%Y"),
    "mode": mode,
    "customer_name": customer_name,
    "customer_address": customer_address,
    "customer_dl": customer_dl,
    "customer_gstin": customer_gstin,
    "customer_pan": customer_pan,
    "taxable_total": taxable_total,
    "sgst_total": sgst_total,
    "cgst_total": cgst_total,
    "gross_total": gross_total,
    "total_qty": total_qty,
    "other": money(other),
    "credit_note": money(credit_note),
    "round_off": round_off,
    "net_amount": net_amount,
    "balance": net_amount,
    "amount_words": amount_in_words(net_amount),
}

# ── Summary ──
st.divider()
m1, m2, m3, m4, m5 = st.columns(5)
m1.metric("Taxable", f"₹{taxable_total:,.2f}")
m2.metric("SGST", f"₹{sgst_total:,.2f}")
m3.metric("CGST", f"₹{cgst_total:,.2f}")
m4.metric("Round Off", f"₹{round_off:,.2f}")
m5.metric("Net Amount", f"₹{net_amount:,.2f}")
st.write("**Amount in Words:**", inv_data["amount_words"])

# ── Downloads ──
st.divider()
pdf_bytes = draw_pdf(inv_data, products)
xlsx_bytes = draw_excel(inv_data, products)

d1, d2 = st.columns(2)
with d1:
    st.download_button(
        "Download PDF", data=pdf_bytes,
        file_name=f"{invoice_no}_tax_invoice.pdf",
        mime="application/pdf", use_container_width=True)
with d2:
    st.download_button(
        "Download Excel", data=xlsx_bytes,
        file_name=f"{invoice_no}_tax_invoice.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True)
